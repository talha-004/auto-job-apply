import os
import re
import threading
from abc import ABC, abstractmethod
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional, Set, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.core.logger import logger
from app.models.job import JobApplicationRecord, ApplicationStatus, ReasonCode


def compute_job_fingerprint(company: str, title: str, location: Optional[str] = None) -> str:
    """
    Computes normalized semantic fingerprint for duplicate detection across URLs.
    Normalizes company suffixes, role abbreviations, and city aliases.
    """
    from app.platforms.naukri_helpers import normalize_token
    norm_comp = normalize_token(company or "")
    suffixes = [
        "private limited", "pvt ltd", "technologies", "technology", 
        "solutions", "services", "corporation", "limited", "corp", 
        "ltd", "inc", "software"
    ]
    for sfx in suffixes:
        if norm_comp.endswith(sfx):
            norm_comp = norm_comp[:-len(sfx)].strip()
            break

    norm_title = normalize_token(title or "")
    norm_title = re.sub(r"\bsr\b", "senior", norm_title)
    norm_title = re.sub(r"\bdev\b", "developer", norm_title)
    norm_title = re.sub(r"\beng\b", "engineer", norm_title)

    norm_loc = normalize_token(location or "")
    norm_loc = re.sub(r"\bbengaluru\b", "bangalore", norm_loc)
    norm_loc = re.sub(r"\bhyd\b", "hyderabad", norm_loc)
    
    # Deduplicate repeated location tokens (e.g. 'bengaluru / bangalore' -> 'bangalore')
    loc_tokens = []
    for tok in norm_loc.split():
        if tok not in loc_tokens:
            loc_tokens.append(tok)
    norm_loc = " ".join(loc_tokens)

    return f"fp:{norm_comp}|{norm_title}|{norm_loc}"


class BaseStorageRepository(ABC):
    """Abstract persistence interface decoupling application runners from concrete storage."""
    @abstractmethod
    def log_application(self, record: JobApplicationRecord) -> bool:
        pass

    @abstractmethod
    def update_application_status(
        self,
        job_url: str,
        new_status: ApplicationStatus,
        notes: Optional[str] = None,
        reason_code: Optional[str] = None
    ) -> bool:
        pass

    @abstractmethod
    def is_already_processed(self, job_url: str) -> bool:
        pass

    @abstractmethod
    def get_all_records(self, limit: int = 200) -> List[Dict[str, Any]]:
        pass

    @abstractmethod
    def get_manual_review_records(self) -> List[Dict[str, Any]]:
        pass


class ExcelTracker(BaseStorageRepository):
    COLUMNS = [
        "Timestamp",
        "Platform",
        "Job Title",
        "Company",
        "Job URL",
        "Status",
        "Match Score",
        "HR Email",
        "Recruiter",
        "Skip Reason",
        "Applied Date",
        "Notes",
        "Job ID",
        "Quality Score",
        "Priority",
        "Reason Code"
    ]

    def __init__(self, file_path: Path = settings.EXCEL_FILE_PATH):
        self.file_path = file_path
        self._lock = threading.Lock()
        self._applied_urls: Set[str] = set()
        self._processed_urls: Dict[str, str] = {}
        self._processed_job_ids: Dict[str, str] = {}
        self._processed_fingerprints: Dict[str, str] = {}
        self._ensure_workbook_exists()
        self._load_existing_urls()

    def _cell(self, row: Any, index: int, default: str = "") -> str:
        """Safely extract and strip a cell value from an Excel row tuple."""
        return str(row[index]).strip() if len(row) > index and row[index] is not None else default

    def _normalize_url(self, url: str) -> str:
        if not url:
            return ""
        return url.strip().split("?")[0].split("#")[0].rstrip("/")

    def _ensure_workbook_exists(self) -> None:
        """Create workbook with styled header if it does not exist."""
        if not self.file_path.exists():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Applications"
            
            # Header styling
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

            ws.append(self.COLUMNS)
            ws.row_dimensions[1].height = 28

            for col_idx in range(1, len(self.COLUMNS) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center

            # Set column widths
            column_widths = {
                1: 20, # Timestamp
                2: 14, # Platform
                3: 30, # Job Title
                4: 25, # Company
                5: 45, # Job URL
                6: 18, # Status
                7: 14, # Match Score
                8: 28, # HR Email
                9: 20, # Recruiter
                10: 25, # Skip Reason
                11: 15, # Applied Date
                12: 35  # Notes
            }
            for col_idx, width in column_widths.items():
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = width

            wb.save(str(self.file_path))
            logger.info(f"Initialized job application tracking workbook at {self.file_path}")

    def _load_existing_urls(self) -> None:
        """Cache all previously applied/discovered job URLs, job IDs, and fingerprints to quickly prevent duplicates."""
        if not self.file_path.exists():
            return
        try:
            wb = openpyxl.load_workbook(str(self.file_path), read_only=True)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not any(row):
                    continue

                raw_url = self._cell(row, 4)
                norm_url = self._normalize_url(raw_url)
                status = self._cell(row, 5)
                job_id = self._cell(row, 12)

                if job_id:
                    self._processed_job_ids[job_id] = status
                if norm_url:
                    self._applied_urls.add(norm_url)
                    self._processed_urls[norm_url] = status

                # Also cache semantic fingerprint if title & company exist
                title = self._cell(row, 2)
                comp = self._cell(row, 3)
                if title and comp:
                    fp = compute_job_fingerprint(comp, title)
                    if fp:
                        self._processed_fingerprints[fp] = status
            wb.close()
            logger.info(f"Loaded {len(self._processed_urls)} tracked URLs, {len(self._processed_job_ids)} job IDs, and {len(self._processed_fingerprints)} fingerprints from Excel.")
        except Exception as e:
            logger.error(f"Error loading existing URLs from Excel: {e}")

    def is_already_processed(self, job_url: str) -> bool:
        """Check if job URL has already been recorded in any state (discovered, skipped, applied)."""
        if not job_url:
            return False
        norm = self._normalize_url(job_url)
        return norm in self._processed_urls

    def get_processed_status(self, job_url: str) -> Optional[str]:
        """Return the prior processing status of the job URL if present."""
        if not job_url:
            return None
        return self._processed_urls.get(self._normalize_url(job_url))

    def is_already_applied(self, job_url: str) -> bool:
        """Check if job URL has already reached an applied or success status."""
        if not job_url:
            return False
        status = self.get_processed_status(job_url)
        if not status:
            return False
        status_lower = status.lower()
        return any(term in status_lower for term in ["success", "applied", "dry run"])

    def check_duplicate(
        self,
        job_url: str,
        company: str,
        title: str = "",
        location: Optional[str] = None,
        job_id: Optional[str] = None,
        job_title: Optional[str] = None
    ) -> Tuple[bool, Optional[str], Optional[str]]:
        """
        Duplicate contract:
        - Job ID match -> definitive duplicate ('EXACT_JOB_ID')
        - URL match -> definitive duplicate ('EXACT_URL')
        - Fingerprint match with different URL -> possible duplicate ('POSSIBLE_DUPLICATE')
        - Unique -> not duplicate
        Returns (is_duplicate, duplicate_type, existing_status)
        """
        effective_title = job_title or title
        # 1. Exact Job ID duplicate (most stable identifier)
        if job_id and job_id in self._processed_job_ids:
            return (True, "EXACT_JOB_ID", self._processed_job_ids[job_id])

        if not job_url:
            return (False, None, None)

        # 2. Exact URL duplicate (definitive)
        if self.is_already_processed(job_url):
            return (True, "EXACT_URL", self.get_processed_status(job_url))

        # 3. Semantic fingerprint duplicate (probabilistic / possible duplicate)
        fp = compute_job_fingerprint(company, effective_title, location)
        if fp and fp in self._processed_fingerprints:
            return (True, "POSSIBLE_DUPLICATE", self._processed_fingerprints[fp])

        return (False, None, None)

    def log_application(self, record: JobApplicationRecord) -> bool:
        """Thread-safely append a new job application record to the Excel sheet. Returns True on success, False on error."""
        with self._lock:
            try:
                wb = openpyxl.load_workbook(str(self.file_path))
                ws = wb.active

                row_values = [
                    record.timestamp,
                    record.platform,
                    record.job_title,
                    record.company,
                    record.job_url,
                    record.status.value if isinstance(record.status, ApplicationStatus) else str(record.status),
                    f"{record.match_score}%" if record.match_score is not None else "",
                    record.hr_email or "",
                    record.recruiter_name or "",
                    record.skip_reason or "",
                    record.applied_date,
                    record.notes,
                    record.job_id or "",
                    f"{record.job_quality_score}%" if record.job_quality_score is not None else "",
                    f"{record.priority_score}%" if record.priority_score is not None else "",
                    record.reason_code or ""
                ]

                ws.append(row_values)
                new_row_idx = ws.max_row
                ws.row_dimensions[new_row_idx].height = 20

                # Style row based on status
                fill_color = None
                status_str = str(record.status)
                if ApplicationStatus.SUCCESS.value in status_str or "Success" in status_str:
                    fill_color = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # light green
                elif ApplicationStatus.FAILED.value in status_str or "Failed" in status_str:
                    fill_color = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # light red
                elif ApplicationStatus.MANUAL_REVIEW_NEEDED.value in status_str or "Manual" in status_str:
                    fill_color = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # light yellow
                elif ApplicationStatus.POSSIBLE_DUPLICATE.value in status_str or "Duplicate" in status_str:
                    fill_color = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid") # light orange
                elif ApplicationStatus.DISCOVERED.value in status_str:
                    fill_color = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid") # light blue
                elif ApplicationStatus.SKIPPED.value in status_str:
                    fill_color = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid") # light slate
                elif ApplicationStatus.APPLYING.value in status_str or "Applying" in status_str:
                    fill_color = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid") # light warm yellow

                border_thin = Border(
                    left=Side(style='thin', color='E2E8F0'),
                    right=Side(style='thin', color='E2E8F0'),
                    top=Side(style='thin', color='E2E8F0'),
                    bottom=Side(style='thin', color='E2E8F0')
                )

                for col_idx in range(1, len(row_values) + 1):
                    cell = ws.cell(row=new_row_idx, column=col_idx)
                    cell.font = Font(name="Calibri", size=10)
                    cell.border = border_thin
                    cell.alignment = Alignment(vertical="center")
                    if fill_color:
                        cell.fill = fill_color

                wb.save(str(self.file_path))
                status_val = str(record.status.value if isinstance(record.status, ApplicationStatus) else record.status)
                if record.job_id:
                    self._processed_job_ids[record.job_id] = status_val
                if record.job_url:
                    norm = self._normalize_url(record.job_url)
                    self._applied_urls.add(norm)
                    self._processed_urls[norm] = status_val
                    fp = compute_job_fingerprint(record.company, record.job_title)
                    if fp:
                        self._processed_fingerprints[fp] = status_val

                logger.info(f"Recorded application to Excel: {record.job_title} at {record.company} [{record.status}]")
                return True
            except Exception as e:
                logger.error(f"Failed to log application to Excel: {e}")
                return False

    def update_application_status(
        self,
        job_url: str,
        new_status: ApplicationStatus,
        notes: Optional[str] = None,
        reason_code: Optional[str] = None
    ) -> bool:
        """Update the status and notes of an existing record matched by job URL."""
        if not job_url:
            return False
        norm_target = self._normalize_url(job_url)

        with self._lock:
            try:
                wb = openpyxl.load_workbook(str(self.file_path))
                ws = wb.active
                target_row = None

                for row_idx in range(2, ws.max_row + 1):
                    cell_val = ws.cell(row=row_idx, column=5).value # Job URL is col 5
                    if cell_val and self._normalize_url(str(cell_val)) == norm_target:
                        target_row = row_idx
                        break

                if not target_row:
                    logger.warning(f"Could not find existing Excel row for URL: {job_url} to update status.")
                    return False

                # Column 6: Status
                status_str = new_status.value if isinstance(new_status, ApplicationStatus) else str(new_status)
                ws.cell(row=target_row, column=6).value = status_str

                # Column 11: Applied Date if now Success/Applied
                if "Success" in status_str or "Applied" in status_str:
                    ws.cell(row=target_row, column=11).value = datetime.now().strftime("%Y-%m-%d")

                # Column 12: Notes
                if notes:
                    ws.cell(row=target_row, column=12).value = notes

                # Column 16: Reason Code if available
                if reason_code and ws.max_column >= 16:
                    ws.cell(row=target_row, column=16).value = str(reason_code)

                # Update fill color
                fill_color = None
                if ApplicationStatus.SUCCESS.value in status_str or "Success" in status_str:
                    fill_color = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                elif ApplicationStatus.FAILED.value in status_str or "Failed" in status_str:
                    fill_color = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                elif ApplicationStatus.MANUAL_REVIEW_NEEDED.value in status_str or "Manual" in status_str:
                    fill_color = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                elif ApplicationStatus.POSSIBLE_DUPLICATE.value in status_str or "Duplicate" in status_str:
                    fill_color = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
                elif ApplicationStatus.APPLYING.value in status_str or "Applying" in status_str:
                    fill_color = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")

                if fill_color:
                    for col_idx in range(1, max(len(self.COLUMNS), ws.max_column) + 1):
                        ws.cell(row=target_row, column=col_idx).fill = fill_color

                wb.save(str(self.file_path))
                self._processed_urls[norm_target] = status_str
                logger.info(f"Updated application status in Excel for {norm_target} -> {status_str}")
                return True
            except Exception as e:
                logger.error(f"Failed to update application status in Excel: {e}")
                return False

    def get_all_records(self, limit: int = 200) -> List[Dict[str, Any]]:
        """Retrieve application records as a list of dictionaries for API/UI."""
        records = []
        if not self.file_path.exists():
            return records

        with self._lock:
            try:
                wb = openpyxl.load_workbook(str(self.file_path), read_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(min_row=2, values_only=True))
                wb.close()

                # Return recent first
                for row in reversed(rows):
                    if not row or not any(row):
                        continue

                    job_url = self._cell(row, 4)
                    job_title = self._cell(row, 2)
                    job_id = self._cell(row, 12)

                    # Diagnostic check: skip completely empty or malformed rows missing title and identifiers
                    if not job_title and not job_url and not job_id:
                        logger.debug("Skipping malformed or unidentifiable row in Excel tracking sheet.")
                        continue

                    rec = {
                        "timestamp": self._cell(row, 0),
                        "platform": self._cell(row, 1),
                        "job_title": job_title,
                        "company": self._cell(row, 3),
                        "job_url": job_url,
                        "status": self._cell(row, 5),
                        "match_score": self._cell(row, 6),
                        "hr_email": self._cell(row, 7),
                        "recruiter_name": self._cell(row, 8),
                        "skip_reason": self._cell(row, 9),
                        "applied_date": self._cell(row, 10),
                        "notes": self._cell(row, 11),
                        "job_id": job_id,
                        "job_quality_score": self._cell(row, 13),
                        "priority_score": self._cell(row, 14),
                        "reason_code": self._cell(row, 15),
                    }
                    records.append(rec)
                    if len(records) >= limit:
                        break
            except Exception as e:
                logger.error(f"Error reading applications from Excel: {e}")

        return records

    def get_manual_review_records(self) -> List[Dict[str, Any]]:
        """Retrieve all records requiring manual review or flagged as possible duplicates."""
        all_records = self.get_all_records(limit=1000)
        review_records = []
        for r in all_records:
            st = str(r.get("status", "")).lower()
            code = str(r.get("reason_code", "")).upper()
            if (
                "manual" in st
                or "duplicate" in st
                or code in ("SUBMISSION_UNKNOWN", "POSSIBLE_DUPLICATE", "JOB_RISK_FLAG", "UNKNOWN_APPLICATION_TYPE", "PROFILE_INCOMPLETE")
            ):
                review_records.append(r)
        return review_records

    def get_statistics(self) -> Dict[str, Any]:
        """Compute summary statistics for dashboard analytics."""
        records = self.get_all_records(limit=2000)
        total = len(records)
        success = 0
        failed = 0
        manual = 0
        platforms: Dict[str, int] = {}

        for r in records:
            st = str(r.get("status", ""))
            plat = str(r.get("platform", "Unknown"))
            platforms[plat] = platforms.get(plat, 0) + 1

            if "Success" in st or "Applied" in st or "Dry Run" in st:
                success += 1
            elif "Failed" in st:
                failed += 1
            elif "Manual" in st or "Duplicate" in st:
                manual += 1

        return {
            "total_applications": total,
            "success_count": success,
            "failed_count": failed,
            "manual_review_count": manual,
            "platform_breakdown": platforms
        }

excel_tracker = ExcelTracker()
